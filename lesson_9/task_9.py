"""
Prosty arkusz kalkulacyjny:
Używając openpyxl, stwórz plik finanse.xlsx.
W kolumnie A umieść nazwy wydatków, w kolumnie B ich wartości.
Pod spodem wstaw formułę =SUM(B1:B2).
"""

from openpyxl import Workbook

wb = Workbook()
ws = wb.active

ws["A1"] = "Czynsz"
ws["B1"] = 2000

ws["A2"] = "Jedzenie"
ws["B2"] = 800

ws["B3"] = "=SUM(B1:B2)"

wb.save("finanse.xlsx")
